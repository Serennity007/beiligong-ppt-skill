# -*- coding: utf-8 -*-
"""v60 步骤1：全篇字号分级放大 + 文字颜色加深 + 收紧行距，并输出逐形状字号映射表"""
import json, os, sys
from pptx import Presentation
from pptx.util import Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE_TYPE

sys.stdout.reconfigure(encoding="utf-8")
B = r"C:/Users/ztl/Desktop/智能体安全_党校授课_交付_20261021/教师意见_五章结构重构_20261017/v59_教师录音纠偏_20261020"
SRC = os.path.join(B, "_build", "底稿_v60_放大字号前.pptx")
OUT = os.path.join(B, "_build", "v60_step1.pptx")
MAP = os.path.join(B, "_build", "size_map.json")


def bump(pt):
    if pt is None:
        return 18.0
    if pt < 11:
        v = max(14.0, pt * 1.40)
    elif pt < 14:
        v = max(16.0, pt * 1.25)
    elif pt < 18:
        v = max(19.0, pt * 1.20)
    elif pt < 22:
        v = max(23.0, pt * 1.18)
    elif pt < 26:
        v = max(27.0, pt * 1.13)
    else:
        v = max(34.0, pt * 1.10)
    return round(v * 2) / 2


CMAP = {
    "006633": "173404", "006C39": "173404", "044A2E": "122C03",
    "17865A": "173404", "4AA56C": "173404", "34726A": "122C03",
    "595959": "2E2E2E", "6E6E6E": "3D3D3D", "222222": "141414",
    "404040": "262626", "4A4A4A": "2A2A2A", "9A9A9A": "555555",
    "BFBFBF": "666666", "C0C0C0": "666666", "CC3333": "B02020",
}
FMAP = {"006633": "173404", "006C39": "173404"}

prs = Presentation(SRC)
IN = 914400.0
mapping = []
stats = {"runs": 0, "shapes": 0}


def walk(shapes, page):
    for sh in shapes:
        if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
            walk(sh.shapes, page)
            continue
        try:
            f = sh.fill
            if f.type is not None and str(f.type).startswith("SOLID"):
                c = str(f.fore_color.rgb).upper()
                if c in FMAP and (sh.width or 0) / IN >= 0.4 and (sh.height or 0) / IN >= 0.28:
                    f.fore_color.rgb = RGBColor.from_string(FMAP[c])
        except Exception:
            pass
        if not sh.has_text_frame:
            continue
        tf = sh.text_frame
        # 逐段收集原字号 -> 目标字号
        pinfo = []
        for p in tf.paragraphs:
            rs = [r for r in p.runs if r.text.strip()]
            if not rs:
                continue
            cur = [(r.font.size.pt if r.font.size else None) for r in rs]
            base = max([c for c in cur if c] or [18.0])
            pinfo.append({"runs": rs, "cur": cur, "orig": base, "new": bump(base)})
        if not pinfo:
            continue
        stats["shapes"] += 1
        for pi in pinfo:
            for r in pi["runs"]:
                stats["runs"] += 1
                r.font.size = Pt(pi["new"])
                try:
                    c = str(r.font.color.rgb).upper()
                except Exception:
                    c = None
                if c in CMAP:
                    r.font.color.rgb = RGBColor.from_string(CMAP[c])
        for p in tf.paragraphs:
            try:
                if p.line_spacing is not None and p.line_spacing > 1.05:
                    p.line_spacing = 1.0
            except Exception:
                pass
            for attr in ("space_before", "space_after"):
                try:
                    v = getattr(p, attr)
                    if v and v > Pt(0):
                        setattr(p, attr, Pt(0))
                except Exception:
                    pass
        mapping.append({
            "page": page, "id": sh.shape_id, "name": sh.name,
            "orig": [pi["orig"] for pi in pinfo],
            "new": [pi["new"] for pi in pinfo],
        })


for i, sl in enumerate(prs.slides, 1):
    walk(sl.shapes, i)

prs.save(OUT)
json.dump(mapping, open(MAP, "w", encoding="utf-8"), ensure_ascii=False)
print("saved", OUT)
print("文字形状", stats["shapes"], "run", stats["runs"])
