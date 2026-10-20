# -*- coding: utf-8 -*-
"""v59 增量改造：目录重排（参照教师样例）、页内小标题绿色横带、右上角完整院标"""
import copy, sys
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE_TYPE, MSO_SHAPE
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn

sys.stdout.reconfigure(encoding="utf-8")
P = r"C:/Users/ztl/Desktop/智能体安全_党校授课_交付_20261021/教师意见_五章结构重构_20261017/v59_教师录音纠偏_20261020/智能体安全_技术趋势与风险防控_v59_白底深字与标题分层_20261020.pptx"
IN = 914400.0
GREEN = RGBColor(0x00, 0x66, 0x33)
DARK = RGBColor(0x17, 0x34, 0x04)
GRAY = RGBColor(0x4A, 0x4A, 0x4A)
LGRAY = RGBColor(0x6E, 0x6E, 0x6E)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FONT = "微软雅黑"

TOC_ROWS = [
    ("01", "智能体的概念与技术构成", "定义与能力 ／ 政策部署 ／ 运行机制 ／ 协作连接 ／ 工具边界"),
    ("02", "风险版图：四层风险与防护控制", "风险总览 ／ 基础设施 ／ 管理 ／ 协同 ／ 应用接口 ／ 防线落地"),
    ("03", "验证与治理：评测、规则与合规", "任务链回看 ／ 评测基准 ／ 治理规则 ／ 国际工具 ／ 模型风险"),
    ("04", "案例复盘：四类事件与证据链", "典型事件 ／ 越界机理 ／ 公开通报 ／ 证据链复盘"),
    ("05", "可信协作：区块链赋能智能体网络", "架构运行 ／ 信任缺口 ／ 链上链下 ／ 基础设施"),
    ("06", "趋势与行动", "运行环境 ／ 四项安全能力 ／ 失控根因 ／ 单位行动与全篇收束"),
]


def remove_shape(sh):
    el = sh._element
    el.getparent().remove(el)


def set_ea_font(run, name=FONT):
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:latin", "a:ea", "a:cs"):
        e = rPr.find(qn(tag))
        if e is None:
            e = rPr.makeelement(qn(tag), {})
            rPr.append(e)
        e.set("typeface", name)


def add_text(slide, l, t, w, h, lines, anchor=MSO_ANCHOR.TOP, align=PP_ALIGN.LEFT):
    """lines: [(text, size, bold, color, spc)]"""
    box = slide.shapes.add_textbox(Emu(int(l * IN)), Emu(int(t * IN)),
                                   Emu(int(w * IN)), Emu(int(h * IN)))
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, (txt, size, bold, color, spc) in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        r = p.add_run()
        r.text = txt
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.color.rgb = color
        set_ea_font(r)
        if spc:
            r._r.get_or_add_rPr().set("spc", str(spc))
    return box


def add_chip(slide, l, t, w, h, text, size=16):
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                Emu(int(l * IN)), Emu(int(t * IN)),
                                Emu(int(w * IN)), Emu(int(h * IN)))
    sh.adjustments[0] = 0.25
    sh.fill.solid()
    sh.fill.fore_color.rgb = GREEN
    sh.line.fill.background()
    sh.shadow.inherit = False
    tf = sh.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = True
    r.font.color.rgb = WHITE
    set_ea_font(r)
    return sh


prs = Presentation(P)
slides = list(prs.slides)
log = []

# ---------- 0. 封面院标模板 ----------
logo_tpl = None
img_part = None
for sh in slides[0].shapes:
    if sh.shape_type == MSO_SHAPE_TYPE.PICTURE:
        logo_tpl = copy.deepcopy(sh._element)
        blip = sh._element.blipFill.find(qn("a:blip"))
        img_part = slides[0].part.related_part(blip.get(qn("r:embed")))
        break
assert logo_tpl is not None, "封面未找到院标图片"

LOGO_H = 0.46
LOGO_W = LOGO_H * 367.0 / 267.0
LOGO_L = 13.333 - 0.42 - LOGO_W
LOGO_T = 0.27


def place_logo(slide, idx, nid):
    # 删旧右上角小图
    for sh in list(slide.shapes):
        if sh.shape_type == MSO_SHAPE_TYPE.PICTURE and sh.top is not None:
            t, l = sh.top / IN, sh.left / IN
            w, h = sh.width / IN, sh.height / IN
            if t < 1.0 and l > 10.5 and w < 2.5 and h < 1.0:
                remove_shape(sh)
                log.append((idx, "删旧院标", sh.name))
    el = copy.deepcopy(logo_tpl)
    blip = el.blipFill.find(qn("a:blip"))
    rid = slide.part.relate_to(img_part, RT.IMAGE)
    blip.set(qn("r:embed"), rid)
    cnv = el.find(qn("p:nvPicPr")).find(qn("p:cNvPr"))
    cnv.set("id", str(nid))
    cnv.set("name", "v59b:logo")
    xfrm = el.find(qn("p:spPr")).find(qn("a:xfrm"))
    off, ext = xfrm.find(qn("a:off")), xfrm.find(qn("a:ext"))
    off.set("x", str(int(LOGO_L * IN)))
    off.set("y", str(int(LOGO_T * IN)))
    ext.set("cx", str(int(LOGO_W * IN)))
    ext.set("cy", str(int(LOGO_H * IN)))
    slide.shapes._spTree.append(el)
    log.append((idx, "加院标", f"{LOGO_W:.2f}x{LOGO_H:.2f}"))


# ---------- 1. 目录页重排 ----------
sl2 = slides[1]
for sh in list(sl2.shapes):
    txt = sh.text_frame.text.strip() if sh.has_text_frame else ""
    l = (sh.left or 0) / IN
    t = (sh.top or 0) / IN
    w = (sh.width or 0) / IN
    drop = False
    if sh.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE and w > 12 and t < 0.5:
        drop = True                                    # 顶部横幅
    elif sh.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE and 0.15 < w < 0.4 and t < 0.5:
        drop = True                                    # 左侧色条
    elif sh.has_text_frame and txt == "报告提纲":
        drop = True                                    # 横幅标题
    elif sh.has_text_frame and "／" in txt and 0.5 < l < 2.0:
        drop = True                                    # 旧目录条目
    if drop:
        remove_shape(sh)
        log.append((2, "目录重排-删", txt[:20] or sh.name))

# 左侧标签
add_text(sl2, 1.15, 2.62, 3.0, 1.0, [("目  录", 44, True, DARK, 0)])
add_text(sl2, 1.18, 3.58, 3.0, 0.45, [("CONTENT", 15, True, LGRAY, 400)])
bar = sl2.shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(int(0.92 * IN)), Emu(int(2.70 * IN)),
                           Emu(int(0.11 * IN)), Emu(int(1.22 * IN)))
bar.fill.solid()
bar.fill.fore_color.rgb = GREEN
bar.line.fill.background()
bar.shadow.inherit = False

# 右侧条目
y = 1.62
for num, title, subs in TOC_ROWS:
    add_chip(sl2, 4.45, y + 0.04, 0.62, 0.40, num, 15)
    add_text(sl2, 5.32, y, 7.5, 0.42, [(title, 20, True, DARK, 0)])
    add_text(sl2, 5.32, y + 0.44, 7.6, 0.34, [(subs, 12, False, GRAY, 0)])
    y += 0.86
log.append((2, "目录重排-建", "目录/CONTENT + 6条目"))

# ---------- 2. 页内小标题绿色横带 ----------
band_cnt = 0
for idx, sl in enumerate(slides, 1):
    if idx in (1, 2, 186):
        continue
    for sh in list(sl.shapes):
        if sh.shape_type != MSO_SHAPE_TYPE.TEXT_BOX or not sh.has_text_frame:
            continue
        t = (sh.top or 0) / IN
        l = (sh.left or 0) / IN
        w = (sh.width or 0) / IN
        txt = sh.text_frame.text.strip()
        if not (0.88 <= t <= 1.18 and l < 1.6 and w > 9 and txt):
            continue
        band = sl.shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(int(0.40 * IN)),
                                   Emu(int(0.96 * IN)), Emu(int(12.53 * IN)),
                                   Emu(int(0.44 * IN)))
        band.fill.solid()
        band.fill.fore_color.rgb = GREEN
        band.line.fill.background()
        band.shadow.inherit = False
        sh._element.addprevious(band._element)
        sh.left = Emu(int(0.68 * IN))
        sh.top = Emu(int(0.96 * IN))
        sh.width = Emu(int(12.0 * IN))
        sh.height = Emu(int(0.44 * IN))
        tf = sh.text_frame
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        for p in tf.paragraphs:
            for r in p.runs:
                r.font.size = Pt(15)
                r.font.bold = True
                r.font.color.rgb = WHITE
        band_cnt += 1
        log.append((idx, "小标题横带", txt[:24]))
        break
print("小标题横带:", band_cnt)

# ---------- 3. 右上角完整院标 ----------
nid = 99000
for idx, sl in enumerate(slides, 1):
    if idx in (1, 186):
        continue
    nid += 1
    place_logo(sl, idx, nid)

prs.save(P)
print("saved:", P)
from collections import Counter
print(Counter(k for _, k, _ in log))
