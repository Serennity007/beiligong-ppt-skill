# -*- coding: utf-8 -*-
"""v59：依据教师录音纠偏 —— 标题层级统一 / 底部与底块改白底 / 白字浅字改深色
范围：不含院标修复，不含总图与目录重构，不含案例改写，不含增删页。
"""
import json, os, re, shutil, copy
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.oxml.ns import qn

BASE = r"C:/Users/ztl/Desktop/智能体安全_党校授课_交付_20261021/教师意见_五章结构重构_20261017"
SRC = os.path.join(BASE, "v58_图文与标题精修_20261020", "智能体安全_技术趋势与风险防控_v58_20261020.pptx")
DST_DIR = os.path.join(BASE, "v59_教师录音纠偏_20261020")
BUILD = os.path.join(DST_DIR, "_build")
DST = os.path.join(DST_DIR, "智能体安全_技术趋势与风险防控_v59_白底深字与标题分层_20261020.pptx")

IN_W = 914400.0
DARK = RGBColor(0x17, 0x34, 0x04)      # 深绿主字色
DARK2 = RGBColor(0x1F, 0x1F, 0x1F)     # 近黑
BODY = RGBColor(0x4A, 0x4A, 0x4A)      # 深灰
LINE_G = RGBColor(0x00, 0x66, 0x33)    # 细线深绿
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
PAGE_NO = RGBColor(0x6E, 0x6E, 0x6E)

GREEN_SET = {"E8F5E9", "EAF3EE", "D6E8C6", "006633", "006C39", "173404", "17865A",
             "044A2E", "3F5B28", "4AA56C", "639922", "8CAA73", "34726A",
             "D9D9D9", "F1F1EF"}

OVERRIDE = {
    "一 · 定义与能力演进": "智能体的定义与能力演进",
    "一 · 政策部署与目标": "智能体的政策部署与目标",
    "一 · 运行机制与关键技术": "智能体的运行机制与关键技术",
    "一 · 多智能体协作与连接": "多智能体协作与连接",
    "一 · 工具调用与应用边界": "工具调用与应用边界",
    "二 · 风险总览": "四层风险总览",
    "二 · 基础设施层": "基础设施层风险与防护",
    "二 · 管理层": "管理层风险与防护",
    "二 · 协同层": "协同层风险与防护",
    "二 · 应用接口层": "应用接口层风险与防护",
    "二 · 防线落地": "四层防线的落地与验收",
    "三 · 回到任务链": "回到任务链看验证环节",
    "三 · 评测基准": "安全评测基准",
    "三 · 我国治理规则与合规要求": "我国治理规则与合规要求",
    "三 · 国际政策工具与标准": "国际政策工具与标准",
    "三 · 模型自身风险": "模型自身风险",
    "四 · 四类典型事件": "四类典型安全事件",
    "四 · 越界机理": "越界机理",
    "四 · 公开通报与近期事件": "公开通报与近期事件",
    "四 · 证据链与复盘方法": "证据链与复盘方法",
    "五 · 协作网络的架构与运行": "协作网络的架构与运行",
    "五 · 跨主体协作的信任缺口与瓶颈": "跨主体协作的信任缺口与瓶颈",
    "五 · 链上机制与链下控制": "链上机制与链下控制",
    "五 · 可信协作基础设施与研究方向": "可信协作基础设施与研究方向",
    "六 · 运行环境": "运行环境与安全要求",
    "六 · 四项安全能力": "四项安全能力",
    "六 · 失控事件与根因": "失控事件与根因",
    "六 · 单位行动与全篇收束": "单位行动与全篇收束",
}
SELF_TITLED = {"总览", "示例", "部分收束"}

log = []


# ---------------- helpers ----------------
def inches(v):
    return None if v is None else v / IN_W


def iter_all(shapes):
    """递归遍历，含组合内形状"""
    for sh in shapes:
        yield sh
        if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
            for s in iter_all(sh.shapes):
                yield s


def remove_shape(sh):
    el = sh._element
    el.getparent().remove(el)


def fill_color_hex(sh):
    try:
        f = sh.fill
        if f.type is not None and str(f.type).startswith("SOLID"):
            return str(f.fore_color.rgb).upper()
    except Exception:
        pass
    return None


def set_fill(sh, rgb):
    try:
        sh.fill.solid()
        sh.fill.fore_color.rgb = rgb
    except Exception:
        pass


def set_no_fill(sh):
    try:
        sh.fill.background()
    except Exception:
        pass


def set_line(sh, rgb, pt):
    try:
        sh.line.color.rgb = rgb
        sh.line.width = Pt(pt)
    except Exception:
        pass


def no_line(sh):
    try:
        sh.line.fill.background()
    except Exception:
        pass


def luminance(rgb_str):
    try:
        r = int(rgb_str[0:2], 16)
        g = int(rgb_str[2:4], 16)
        b = int(rgb_str[4:6], 16)
    except Exception:
        return 1.0
    return (0.299 * r + 0.587 * g + 0.114 * b) / 255.0


def set_single_line(tf, text, size, color, bold, anchor=MSO_ANCHOR.TOP):
    tf.word_wrap = True
    paras = list(tf.paragraphs)
    for p in paras[1:]:
        try:
            p._p.getparent().remove(p._p)
        except Exception:
            pass
    p0 = tf.paragraphs[0]
    runs = list(p0.runs)
    if runs:
        runs[0].text = text
        for r in runs[1:]:
            try:
                r._r.getparent().remove(r._r)
            except Exception:
                pass
        run = runs[0]
    else:
        run = p0.add_run()
        run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    p0.alignment = PP_ALIGN.LEFT
    try:
        tf.vertical_anchor = anchor
    except Exception:
        pass


def rec(idx, kind, detail):
    log.append({"page": idx, "kind": kind, "detail": detail})


# ---------------- 1. 先算出每页的统摄标题 ----------------
prs = Presentation(SRC)
slides = list(prs.slides)

# ---- 0. 页面背景：浅绿铺底 → 白底 ----
for i, sl in enumerate(slides, 1):
    try:
        f = sl.background.fill
        f.solid()
        f.fore_color.rgb = WHITE
        rec(i, "页面背景", "E8F5E9 → 白底")
    except Exception as e:
        rec(i, "页面背景失败", str(e))

page_tag, page_title = {}, {}
for i, sl in enumerate(slides, 1):
    tag = None
    title = None
    for sh in iter_all(sl.shapes):
        if sh.name == "v57:nodeTag" and sh.has_text_frame and sh.text_frame.text.strip():
            tag = sh.text_frame.text.strip()
        if (sh.has_text_frame and sh.text_frame.text.strip()
                and sh.left is not None and inches(sh.left) < 5
                and sh.top is not None and inches(sh.top) < 0.7
                and sh.height is not None and inches(sh.height) < 1.0
                and sh.width is not None and inches(sh.width) > 3
                and sh.name != "chrome:pageno"):
            title = sh.text_frame.text.strip()
    page_tag[i] = tag
    page_title[i] = title

groups = {}
for i in sorted(page_tag):
    t = page_tag[i]
    if t:
        groups.setdefault(t, []).append(i)

gov_of, sub_of = {}, {}
for tag, pages in groups.items():
    m = re.match(r"^([一二三四五六])\s*·\s*(.+)$", tag)
    node = m.group(2) if m else tag
    if len(pages) >= 2 and node not in SELF_TITLED:
        gov = OVERRIDE.get(tag, node)
        for p in pages:
            gov_of[p] = gov
            sub_of[p] = page_title.get(p) or ""
    else:
        for p in pages:
            gov_of[p] = page_title.get(p) or node
            sub_of[p] = None
for i in sorted(page_tag):
    if not page_tag[i] and page_title.get(i):
        gov_of[i] = page_title[i]
        sub_of[i] = None

# ---------------- 2. 逐页改造 ----------------
SECTION_PAGES = {3, 31, 61, 84, 122, 164}
COVER_PAGES = {1, 186}

for idx, sl in enumerate(slides, 1):
    top = list(sl.shapes)

    # ---- 部分隔页 ----
    if idx in SECTION_PAGES:
        for sh in list(sl.shapes):
            if sh.name == "v57:bg":
                set_fill(sh, WHITE)
                no_line(sh)
                rec(idx, "部分隔页背景", "E8F5E9 → 白底")
            elif sh.name == "v57:oneliner":
                remove_shape(sh)
                rec(idx, "删导读", "删除“本部分回答：……”")
            elif sh.name and sh.name.startswith("v57:range"):
                remove_shape(sh)
                rec(idx, "删页段", f"删除页码索引 {sh.name}")
        continue

    if idx in COVER_PAGES:
        continue

    # ---- 定位 chrome 要素 ----
    topbar = None
    accent = None
    titlebox = None
    sub_band = None
    sub_box = None
    bottom_bars = []
    nodetag = None
    for sh in top:
        if sh.shape_type not in (MSO_SHAPE_TYPE.AUTO_SHAPE, MSO_SHAPE_TYPE.TEXT_BOX,
                                 MSO_SHAPE_TYPE.FREEFORM):
            continue
        l, t = inches(sh.left), inches(sh.top)
        w, h = inches(sh.width), inches(sh.height)
        if l is None or t is None or w is None or h is None:
            continue
        if sh.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE:
            if w > 12 and h < 1.2 and t < 0.5:
                topbar = sh
            elif 0.15 < w < 0.4 and 0.5 < h < 1.0 and t < 0.5:
                accent = sh
            elif w > 12 and 0.85 < t < 1.25 and 0.4 < h < 1.35:
                sub_band = sh
            elif w > 8 and h < 1.3 and t > 5.6:
                bottom_bars.append(sh)
        else:
            if sh.name == "v57:nodeTag":
                nodetag = sh
            elif t < 0.7 and l < 5 and h < 1.0 and w > 3 and sh.has_text_frame \
                    and sh.text_frame.text.strip():
                titlebox = sh
            elif 0.85 < t < 1.25 and l < 1.5 and w > 8 and sh.has_text_frame \
                    and sh.text_frame.text.strip():
                sub_box = sh

    # ---- A. 顶部横幅：白底 + 深绿细线 + 深色统摄标题 ----
    if topbar is not None:
        old = fill_color_hex(topbar)
        set_fill(topbar, WHITE)
        set_line(topbar, LINE_G, 1.25)
        rec(idx, "顶部横幅", f"{old} → 白底 + 深绿细线")

    gov = gov_of.get(idx)
    if titlebox is not None and gov:
        old_text = titlebox.text_frame.text.strip()
        old_size = None
        for p in titlebox.text_frame.paragraphs:
            for r in p.runs:
                if r.font.size:
                    old_size = r.font.size.pt
                    break
        size = old_size if (old_size and 24 <= old_size <= 34) else 28
        set_single_line(titlebox.text_frame, gov, size, DARK, True, MSO_ANCHOR.MIDDLE)
        set_no_fill(titlebox)
        # 右上角主题标签删除后，标题可用宽度放宽到院标左侧
        if inches(titlebox.width) < 11.3:
            titlebox.width = Emu(int(11.3 * IN_W))
        if inches(titlebox.height) < 0.6:
            titlebox.height = Emu(int(0.64 * IN_W))
        if old_text != gov:
            rec(idx, "统摄标题", f"“{old_text}” → “{gov}”")

    if nodetag is not None:
        txt = nodetag.text_frame.text.strip()
        remove_shape(nodetag)
        rec(idx, "删右上角缩略主题", f"删除 “{txt}”")

    # ---- B. 页内小标题（原副标题位）----
    sub = sub_of.get(idx)
    if sub_box is not None:
        if sub:
            set_single_line(sub_box.text_frame, sub, 18, DARK, True, MSO_ANCHOR.TOP)
            set_no_fill(sub_box)
            if inches(sub_box.height) > 0.55:
                sub_box.height = Emu(int(0.45 * IN_W))
            rec(idx, "页内小标题", f"“{sub}”")
        else:
            remove_shape(sub_box)
            rec(idx, "删固定副标题", "统摄标题即本页论题，副标题句删除")
    elif sub:
        rec(idx, "缺页内小标题位", f"未找到副标题框，小标题“{sub}”未落地")

    if sub_band is not None:
        if sub:
            set_no_fill(sub_band)
            no_line(sub_band)
            rec(idx, "副标题底带", "灰/绿底 → 无填充")
        else:
            remove_shape(sub_band)
            rec(idx, "删副标题底带", "无小标题时整块删除")

    # ---- C. 底部横幅：白底 + 深绿细线 ----
    for b in bottom_bars:
        old = fill_color_hex(b)
        set_fill(b, WHITE)
        set_line(b, LINE_G, 1.0)
        rec(idx, "底部横幅", f"{old} → 白底 + 深绿细线")

# 目录页（第2页）删导读性副标题与总结底栏
sl2 = slides[1]
for sh in list(sl2.shapes):
    if sh.name in ("chrome:subtitle", "chrome:takeaway", "chrome:takeaway-bar", "chrome:band"):
        nm = sh.name
        remove_shape(sh)
        rec(2, "目录页删导读", f"删除 {nm}")

# ---------------- 3. 全篇去绿底 / 灰底 ----------------
for idx, sl in enumerate(slides, 1):
    if idx in COVER_PAGES:
        continue
    for sh in iter_all(sl.shapes):
        if sh.shape_type not in (MSO_SHAPE_TYPE.AUTO_SHAPE, MSO_SHAPE_TYPE.TEXT_BOX,
                                 MSO_SHAPE_TYPE.FREEFORM):
            continue
        c = fill_color_hex(sh)
        if not c or c not in GREEN_SET:
            continue
        w = inches(sh.width) or 0
        h = inches(sh.height) or 0
        if w > 12.5 and h > 6.9:              # 整页背景
            set_fill(sh, WHITE)
            no_line(sh)
            rec(idx, "整页底", f"{c} → 白底")
        elif w >= 1.8 and h >= 0.28:          # 卡片 / 底块
            set_fill(sh, WHITE)
            set_line(sh, LINE_G, 0.75)
            rec(idx, "卡片去绿底", f"{c} {round(w,2)}×{round(h,2)} → 白底 + 细线")
        # 小面积重点标识（左色条、圆点、细线）保留原色

# ---------------- 4. 字体颜色加深 ----------------
for idx, sl in enumerate(slides, 1):
    if idx in COVER_PAGES:
        continue
    for sh in iter_all(sl.shapes):
        if not sh.has_text_frame:
            continue
        c = fill_color_hex(sh)
        light = True if c is None else luminance(c) > 0.62
        for p in sh.text_frame.paragraphs:
            for r in p.runs:
                try:
                    col = str(r.font.color.rgb).upper()
                except Exception:
                    col = None
                if col == "FFFFFF" and light:
                    r.font.color.rgb = DARK
                    rec(idx, "白字转深", r.text[:24])
                elif col in ("9A9A9A", "BFBFBF", "C0C0C0"):
                    r.font.color.rgb = PAGE_NO
                    rec(idx, "浅灰字转深", r.text[:24])

prs.save(DST)
with open(os.path.join(BUILD, "v59_changes.json"), "w", encoding="utf-8") as f:
    json.dump(log, f, ensure_ascii=False, indent=1)

from collections import Counter
cnt = Counter(x["kind"] for x in log)
print("saved:", DST)
for k, v in cnt.most_common():
    print(f"  {v:5d}  {k}")
print("总改动:", len(log))
