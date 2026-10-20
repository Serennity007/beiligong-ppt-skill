# -*- coding: utf-8 -*-
"""v60 步骤5：恢复 P157 被误改的字号（以实测值为准），仅对真正溢出的标签缩小"""
import json, sys
from pptx import Presentation
from pptx.util import Pt
from pptx.enum.shapes import MSO_SHAPE_TYPE

sys.stdout.reconfigure(encoding="utf-8")
P = r"C:/Users/ztl/Desktop/智能体安全_党校授课_交付_20261021/教师意见_五章结构重构_20261017/v59_教师录音纠偏_20261020/智能体安全_技术趋势与风险防控_v60__20261020.pptx"
B = r"C:/Users/ztl/Desktop/智能体安全_党校授课_交付_20261021/教师意见_五章结构重构_20261017/v59_教师录音纠偏_20261020/_build"
IN = 914400.0

fit = json.load(open(B + "/fit_result.json", encoding="utf-8"))
p157 = {r["id"]: r for r in fit if r["page"] == 157}
bad = {r["id"]: r for r in p157.values() if r["over"] > 3.0}

prs = Presentation(P)
restored, shrunk = 0, 0


def walk(shapes):
    global restored, shrunk
    for sh in shapes:
        if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
            walk(sh.shapes)
            continue
        if not sh.has_text_frame or sh.shape_id not in p157:
            continue
        r = p157[sh.shape_id]
        applied = r["applied"]
        paras = [p for p in sh.text_frame.paragraphs
                 if any(x.text.strip() for x in p.runs)]
        for p, v in zip(paras, applied):
            for x in p.runs:
                if x.text.strip():
                    x.font.size = Pt(v)
        restored += 1
        if sh.shape_id in bad:
            w_in = (sh.width or 0) / IN
            h_in = (sh.height or 0) / IN
            if w_in <= 0.2:
                return0 = None
            else:
                newv = []
                for p, v in zip(paras, applied):
                    s = "".join(x.text for x in p.runs)
                    n = max(1, len(s))
                    fs = min(v, 72.0 * w_in / (n * 1.06), h_in * 72 * 1.15)
                    fs = max(8.0, round(fs * 2) / 2)
                    newv.append(fs)
                    for x in p.runs:
                        if x.text.strip():
                            x.font.size = Pt(fs)
                shrunk += 1
                print("  缩小 id", sh.shape_id, newv,
                      "".join(x.text for x in paras[0].runs)[:18])


walk(prs.slides[156].shapes)
prs.save(P)
print("恢复", restored, "缩小", shrunk)
