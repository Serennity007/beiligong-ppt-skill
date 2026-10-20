# -*- coding: utf-8 -*-
"""v59 核验：残留白字/浅绿底/深底白字风险"""
import sys
from collections import Counter
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

sys.stdout.reconfigure(encoding="utf-8")
P = r"C:/Users/ztl/Desktop/智能体安全_党校授课_交付_20261021/教师意见_五章结构重构_20261017/v59_教师录音纠偏_20261020/智能体安全_技术趋势与风险防控_v59_白底深字与标题分层_20261020.pptx"
IN = 914400.0


def lum(h):
    try:
        r, g, b = int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16)
    except Exception:
        return 1.0
    return (0.299 * r + 0.587 * g + 0.114 * b) / 255.0


def iter_all(shapes):
    for sh in shapes:
        yield sh
        if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
            for s in iter_all(sh.shapes):
                yield s


def fill_hex(sh):
    try:
        f = sh.fill
        if f.type is not None and str(f.type).startswith("SOLID"):
            return str(f.fore_color.rgb).upper()
    except Exception:
        pass
    return None


prs = Presentation(P)
white_on_dark = []
left_green = Counter()
small_white = Counter()

for i, sl in enumerate(prs.slides, 1):
    for sh in iter_all(sl.shapes):
        if not sh.has_text_frame:
            continue
        c = fill_hex(sh)
        for p in sh.text_frame.paragraphs:
            for r in p.runs:
                try:
                    col = str(r.font.color.rgb).upper()
                except Exception:
                    col = None
                if not r.text.strip():
                    continue
                if col == "FFFFFF":
                    if c is None or lum(c) > 0.62:
                        white_on_dark.append((i, sh.name, c, r.text[:30]))
                    else:
                        small_white[(c,)] += 1
        c2 = fill_hex(sh)
        if c2 in ("E8F5E9", "D9D9D9", "006633", "006C39"):
            w = (sh.width or 0) / IN
            h = (sh.height or 0) / IN
            if w >= 1.8 and h >= 0.28:
                left_green[c2] += 1

print("残留白字（所在形状为白底/无填充）:", len(white_on_dark))
for x in white_on_dark[:30]:
    print("  ", x)
print("\n仍保留的白字所在深色形状分布:", dict(small_white))
print("\n残留大块绿/灰底:", dict(left_green))
