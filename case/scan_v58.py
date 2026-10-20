# -*- coding: utf-8 -*-
"""扫描 v58 底稿结构：形状、填充、文字颜色"""
import json, sys
from pptx import Presentation
from pptx.util import Emu

SRC = r"C:/Users/ztl/Desktop/智能体安全_党校授课_交付_20261021/教师意见_五章结构重构_20261017/v58_图文与标题精修_20261020/智能体安全_技术趋势与风险防控_v58_20261020.pptx"
OUT = r"C:/Users/ztl/Desktop/智能体安全_党校授课_交付_20261021/教师意见_五章结构重构_20261017/v59_教师录音纠偏_20261020/_build/scan_v58.json"

def emu2in(v):
    return None if v is None else round(v / 914400.0, 3)

def fill_info(shape):
    try:
        f = shape.fill
    except Exception:
        return None
    try:
        t = str(f.type)
    except Exception:
        return None
    info = {"type": t}
    try:
        if f.type is not None and str(f.type).startswith("SOLID"):
            info["fg"] = str(f.fore_color.rgb)
    except Exception:
        pass
    return info

def run_colors(tf):
    out = []
    try:
        for p in tf.paragraphs:
            for r in p.runs:
                c = None
                try:
                    c = str(r.font.color.rgb)
                except Exception:
                    c = None
                out.append({"t": r.text[:40], "color": c,
                            "sz": (r.font.size.pt if r.font.size else None),
                            "b": r.font.bold})
    except Exception:
        pass
    return out

def walk(shapes, depth=0, prefix=""):
    res = []
    for sh in shapes:
        d = {
            "name": sh.shape_type and str(sh.name),
            "type": str(sh.shape_type),
            "id": sh.shape_id,
            "l": emu2in(sh.left), "t": emu2in(sh.top),
            "w": emu2in(sh.width), "h": emu2in(sh.height),
            "fill": fill_info(sh),
        }
        try:
            d["line"] = str(sh.line.color.rgb) if (sh.line.fill.type is not None and str(sh.line.fill.type).startswith("SOLID")) else None
        except Exception:
            d["line"] = None
        if sh.has_text_frame:
            d["text"] = sh.text_frame.text[:200]
            d["runs"] = run_colors(sh.text_frame)
        if sh.shape_type == 6:  # group
            d["children"] = walk(sh.shapes, depth + 1)
        res.append(d)
    return res

def main():
    prs = Presentation(SRC)
    data = {
        "slide_w": emu2in(prs.slide_width),
        "slide_h": emu2in(prs.slide_height),
        "n": len(prs.slides),
        "layout_names": [l.name for l in prs.slide_layouts],
        "slides": [],
    }
    for i, s in enumerate(prs.slides, 1):
        data["slides"].append({
            "idx": i,
            "layout": s.slide_layout.name if s.slide_layout else None,
            "shapes": walk(s.shapes),
        })
    with open(OUT, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=1)
    print("slides:", data["n"], "size:", data["slide_w"], "x", data["slide_h"])

if __name__ == "__main__":
    main()
