# -*- coding: utf-8 -*-
"""v60 补充修复：少数回退后仍溢出的框（百分比框加宽、胶囊/标签缩小到能放下）"""
import sys
from pptx import Presentation
from pptx.util import Pt, Emu
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.dml.color import RGBColor

sys.stdout.reconfigure(encoding="utf-8")
P = r"C:/Users/ztl/Desktop/智能体安全_党校授课_交付_20261021/教师意见_五章结构重构_20261017/v59_教师录音纠偏_20261020/智能体安全_技术趋势与风险防控_v60__20261020.pptx"
IN = 914400.0
prs = Presentation(P)
fixed = []


def walk(shapes, page):
    for sh in shapes:
        if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
            walk(sh.shapes, page)
            continue
        if not sh.has_text_frame:
            continue
        txt = sh.text_frame.text.strip()
        if not txt:
            continue

        # P52 三个百分比框：加宽
        if page == 52 and txt in ("56.6%", "64.2%", "74.9%"):
            sh.width = Emu(int(1.38 * IN))
            fixed.append((page, sh.shape_id, "加宽", txt))

        # P127 中间胶囊文字
        elif page == 127 and txt == "多智能体协作":
            for p in sh.text_frame.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(15)
            fixed.append((page, sh.shape_id, "缩至15pt", txt))

        # P156 六个序号标题：缩到一行能放下
        elif page == 156 and txt[:1] in "①②③④⑤⑥":
            for p in sh.text_frame.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(14)
            fixed.append((page, sh.shape_id, "缩至14pt", txt))

        # P157 小标签：按宽度算出能放下的字号
        elif page == 157:
            w_in = (sh.width or 0) / IN
            h_in = (sh.height or 0) / IN
            paras = [p for p in sh.text_frame.paragraphs
                     if any(r.text.strip() for r in p.runs)]
            if w_in > 0.2 and paras:
                per = []
                ok = True
                for p in paras:
                    s = "".join(r.text for r in p.runs)
                    n = max(1, len(s))
                    fs = 72.0 * w_in / (n * 1.06)
                    fs = min(fs, 16.0, h_in * 72 * 1.15)
                    fs = max(8.0, round(fs * 2) / 2)
                    per.append(fs)
                for p, fs in zip(paras, per):
                    for r in p.runs:
                        r.font.size = Pt(fs)
                fixed.append((page, sh.shape_id, f"缩至{per}", txt[:16]))


for i, sl in enumerate(prs.slides, 1):
    walk(sl.shapes, i)

prs.save(P)
print("修复", len(fixed), "处")
for f in fixed:
    print("  ", f)
