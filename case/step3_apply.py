# -*- coding: utf-8 -*-
"""v60 步骤3：把实测得到的最终字号写回 PPTX（python-pptx 保存，避免 COM 改写）"""
import json, os, sys, zipfile, re
from pptx import Presentation
from pptx.util import Pt
from pptx.enum.shapes import MSO_SHAPE_TYPE
from collections import Counter

sys.stdout.reconfigure(encoding="utf-8")
B = r"C:/Users/ztl/Desktop/智能体安全_党校授课_交付_20261021/教师意见_五章结构重构_20261017/v59_教师录音纠偏_20261020"
SRC = os.path.join(B, "_build", "v60_step1.pptx")
OUT = os.path.join(B, "_build", "v60_字号放大_20261020.pptx")
FINAL = os.path.join(B, "智能体安全_技术趋势与风险防控_v60__20261020.pptx")

fit = json.load(open(os.path.join(B, "_build", "fit_result.json"), encoding="utf-8"))
mp = {(r["page"], r["id"]): r["applied"] for r in fit}

prs = Presentation(SRC)
n_hit = 0
sizes = Counter()


def walk(shapes, page):
    global n_hit
    for sh in shapes:
        if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
            walk(sh.shapes, page)
            continue
        if not sh.has_text_frame:
            continue
        vals = mp.get((page, sh.shape_id))
        if vals is None:
            continue
        pi = 0
        for p in sh.text_frame.paragraphs:
            rs = [r for r in p.runs if r.text.strip()]
            if not rs:
                continue
            v = vals[pi] if pi < len(vals) else vals[-1]
            pi += 1
            for r in rs:
                r.font.size = Pt(v)
                sizes[v] += 1
        n_hit += 1


for i, sl in enumerate(prs.slides, 1):
    walk(sl.shapes, i)

prs.save(OUT)
print("写回形状:", n_hit, "/", len(fit))

z = zipfile.ZipFile(OUT)
print("sldId:", z.read("ppt/presentation.xml").decode("utf-8").count("<p:sldId "))
print("字号分布(前14):")
for k, v in sizes.most_common(14):
    print(f"   {k} : {v}")

# 覆盖交付文件
import shutil
shutil.copyfile(OUT, FINAL)
print("交付文件已更新:", FINAL)
