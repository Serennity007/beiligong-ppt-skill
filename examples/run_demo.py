"""Exercise scan and size mapping with an invented slide, without Office."""
import json
from pathlib import Path
import subprocess
import sys
import tempfile

from pptx import Presentation
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]


def main():
    parent = ROOT / '_build'
    parent.mkdir(exist_ok=True)
    output = Path(tempfile.mkdtemp(prefix='demo-', dir=parent))
    src, enlarged = output / 'source.pptx', output / 'enlarged.pptx'
    scan, mapping = output / 'scan.json', output / 'size_map.json'
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(7), Inches(2))
    box.text = 'Synthetic demo: font-size mapping'
    box.text_frame.paragraphs[0].runs[0].font.size = Pt(14)
    prs.save(src)
    before = src.read_bytes()
    for script, args in [('ppt_scan.py', [src, scan]),
                         ('ppt_enlarge.py', [src, enlarged, '--map', mapping])]:
        subprocess.run([sys.executable, '-X', 'utf8', str(ROOT / 'scripts' / script), *map(str, args)], check=True)
    scanned = json.loads(scan.read_text(encoding='utf-8'))
    mapped = json.loads(mapping.read_text(encoding='utf-8'))
    result = Presentation(enlarged)
    size = result.slides[0].shapes[0].text_frame.paragraphs[0].runs[0].font.size.pt
    if scanned['n'] != 1 or len(result.slides) != 1 or not mapped or size <= 14 or src.read_bytes() != before:
        raise RuntimeError('Synthetic scan/size/source-preservation check failed')
    print(f'PASS: {output}')
    print('Synthetic scan and size mapping only; no Office measurement or visual delivery check.')


if __name__ == '__main__':
    main()
