"""Synthetic regression checks; no original course deck or Office required."""
import json
from pathlib import Path
import runpy
import subprocess
import sys
import tempfile
from types import ModuleType, SimpleNamespace as NS
import unittest
from unittest.mock import patch

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]


class AuditTests(unittest.TestCase):
    def test_enlarge_preserves_inherited_sizes_and_blank_paragraph_indices(self):
        with tempfile.TemporaryDirectory() as folder:
            tmp = Path(folder)
            prs = Presentation()
            sl = prs.slides.add_slide(prs.slide_layouts[6])
            box = sl.shapes.add_textbox(Inches(1), Inches(2), Inches(8), Inches(2))
            box.text = 'Inherited size\n\nSecond paragraph'
            box.text_frame.paragraphs[0].font.size = Pt(40)
            box.text_frame.paragraphs[2].font.size = Pt(20)
            unknown = sl.shapes.add_textbox(Inches(1), Inches(4), Inches(8), Inches(1))
            unknown.text = 'Unknown inherited theme size'
            src, out, mp = tmp / 'in.pptx', tmp / 'out.pptx', tmp / 'map.json'
            prs.save(src)
            result = subprocess.run([sys.executable, str(ROOT / 'scripts/ppt_enlarge.py'),
                                     str(src), str(out), '--map', str(mp)], capture_output=True)
            self.assertEqual(result.returncode, 0, result.stderr)
            deck = Presentation(out)
            self.assertGreaterEqual(deck.slides[0].shapes[0].text_frame.paragraphs[0].runs[0].font.size.pt, 40)
            self.assertIsNone(deck.slides[0].shapes[1].text_frame.paragraphs[0].runs[0].font.size)
            mapping = json.loads(mp.read_text(encoding='utf-8'))
            self.assertEqual(mapping[0]['paragraphs'], [1, 3])
            fit = tmp / 'fit.json'
            fit.write_text(json.dumps([{'page': 1, 'id': box.shape_id, 'paragraphs': [1, 3],
                                       'applied': [40, 20]}]), encoding='utf-8')
            result = subprocess.run([sys.executable, str(ROOT / 'scripts/ppt_apply.py'),
                                     str(out), '--fit', str(fit), '--out', str(tmp / 'final.pptx')], capture_output=True)
            self.assertEqual(result.returncode, 0, result.stderr)
            final = Presentation(tmp / 'final.pptx').slides[0].shapes[0].text_frame.paragraphs
            self.assertEqual(final[1].text, '')
            self.assertEqual(final[2].runs[0].font.size.pt, 20)

    def test_white_preserves_band_contrast_and_body_during_purge(self):
        from pptx.enum.shapes import MSO_SHAPE
        with tempfile.TemporaryDirectory() as folder:
            tmp = Path(folder)
            prs = Presentation()
            sl = prs.slides.add_slide(prs.slide_layouts[6])
            band = sl.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1), Inches(2), Inches(8), Inches(1))
            band.fill.solid()
            band.fill.fore_color.rgb = RGBColor.from_string('006633')
            text = sl.shapes.add_textbox(Inches(1), Inches(2), Inches(8), Inches(1))
            text.text = 'White on dark band'
            text.text_frame.paragraphs[0].runs[0].font.color.rgb = RGBColor(255,255,255)
            body = sl.shapes.add_textbox(Inches(1), Inches(4), Inches(8), Inches(2))
            body.text = 'Important factual content\nTODO: confirm citation\nDiscussion of TODO in a sentence'
            src, out = tmp / 'in.pptx', tmp / 'out.pptx'
            prs.save(src)
            result = subprocess.run([sys.executable, str(ROOT / 'scripts/ppt_white.py'), str(src), str(out),
                                     '--config', str(ROOT / 'case/white_config.json')], capture_output=True)
            self.assertEqual(result.returncode, 0, result.stderr)
            shapes = Presentation(out).slides[0].shapes
            self.assertEqual(str(shapes[1].text_frame.paragraphs[0].runs[0].font.color.rgb), 'FFFFFF')
            self.assertEqual(shapes[2].text, 'Important factual content\nDiscussion of TODO in a sentence')

    def test_crop_font_fill_and_missing_measurement(self):
        with tempfile.TemporaryDirectory() as tmp:
            tmp = Path(tmp)
            prs = Presentation()
            prs.slide_width = Inches(13.333)
            for page in range(2):
                sl = prs.slides.add_slide(prs.slide_layouts[6])
                sl.background.fill.solid()
                sl.background.fill.fore_color.rgb = RGBColor(255, 255, 255)
                box = sl.shapes.add_textbox(Inches(.5), Inches(.1), Inches(8), Inches(.5))
                box.text = 'A continuous teaching topic'
                box.text_frame.paragraphs[0].runs[0].font.size = Pt(20)
                body = sl.shapes.add_textbox(Inches(1), Inches(2), Inches(8), Inches(1))
                body.text = 'Long body text below the minimum readable size'
                body.text_frame.paragraphs[0].runs[0].font.size = Pt(9)
                body.fill.solid()
                body.fill.fore_color.rgb = RGBColor.from_string('F5F5F5')
                pic = sl.shapes.add_picture(str(ROOT / 'assets/logo_示例_人工智能学院.png'),
                                           Inches(12), Inches(.2), height=Inches(.56))
                if page == 1:
                    pic.crop_left = .2
            src, report = tmp / 'sample.pptx', tmp / 'report.json'
            prs.save(src)
            proc = subprocess.run([sys.executable, str(ROOT / 'scripts/ppt_audit.py'),
                                   str(src), '--json', str(report)], capture_output=True)
            self.assertEqual(proc.returncode, 1, proc.stderr)
            checks = {r['id']: r for r in json.loads(report.read_text(encoding='utf-8'))['checks']}
            self.assertEqual(checks['E6']['level'], 'FAIL')
            self.assertTrue(any(r.get('page') == 2 for r in checks['E6']['items']))
            self.assertEqual(checks['F2']['level'], 'FAIL')
            self.assertEqual(checks['E1b']['level'], 'FAIL')
            self.assertEqual(checks['F3']['level'], 'WARN')
            for values, cid in [([], 'F3-data'), ([{'page': 1, 'name': 'box', 'over': 8}], 'F3')]:
                fit = tmp / 'fit.json'
                fit.write_text(json.dumps(values), encoding='utf-8')
                proc = subprocess.run([sys.executable, str(ROOT / 'scripts/ppt_audit.py'),
                                       str(src), '--fit', str(fit), '--json', str(report)],
                                      capture_output=True)
                self.assertEqual(proc.returncode, 1)
                checks = {r['id']: r for r in json.loads(report.read_text(encoding='utf-8'))['checks']}
                self.assertEqual(checks[cid]['level'], 'FAIL')

    def test_logo_uses_source_ratio_and_actual_slide_width(self):
        with tempfile.TemporaryDirectory() as tmp:
            tmp = Path(tmp)
            prs = Presentation()
            prs.slide_width = Inches(10)
            prs.slides.add_slide(prs.slide_layouts[6])
            src, out, cfg = tmp / 'in.pptx', tmp / 'out.pptx', tmp / 'config.json'
            prs.save(src)
            cfg.write_text(json.dumps({'logo': {'on': True, 'h': .56, 'ratio': 3,
                'right': .38, 'image': str(ROOT / 'assets/logo_示例_人工智能学院.png')}}), encoding='utf-8')
            proc = subprocess.run([sys.executable, str(ROOT / 'scripts/ppt_deco.py'),
                                   str(src), str(out), '--config', str(cfg)], capture_output=True)
            self.assertEqual(proc.returncode, 0, proc.stderr)
            pic = Presentation(out).slides[0].shapes[0]
            w, h = pic.image.size
            self.assertAlmostEqual(pic.width / pic.height, w / h, places=4)
            self.assertAlmostEqual((pic.left + pic.width) / Inches(1), 9.62, places=4)


class FitTests(unittest.TestCase):
    def run_fit(self, original, target, steps='1,.94,.88,.82,.76,.70', broken=False, indices=None):
        # Fake COM isolates the rollback algorithm; real rendering is a separate gate.
        pars = [NS(Font=NS(Size=v)) for v in target]
        physical = list(pars)
        if indices:
            physical = [NS(Font=NS(Size=18)) for _ in range(max(indices))]
            for idx, para in zip(indices, pars):
                physical[idx - 1] = para
        class Range:
            Text = 'sample'
            Paragraphs = NS(Count=len(physical), Item=lambda i: physical[i-1])
            BoundLeft = BoundTop = 0
            BoundWidth = 100
            @property
            def BoundHeight(self):
                if broken:
                    raise RuntimeError('measurement failed')
                return 100 + max(p.Font.Size - o for p, o in zip(pars, original))
        shape = NS(Type=1, Id=5, HasTextFrame=True, TextFrame2=NS(TextRange=Range()),
                   Left=0, Top=0, Width=100, Height=100)
        slide = NS(Shapes=NS(Count=1, Item=lambda i: shape))
        class Slides:
            Count = 1
            def __call__(self, i):
                return slide
        deck = NS(Slides=Slides(), Close=lambda: None)
        app = NS(Presentations=NS(Open=lambda *args, **kwargs: deck), Quit=lambda: None)
        win32 = ModuleType('win32com')
        client = ModuleType('win32com.client')
        client.DispatchEx = lambda name: app
        win32.client = client
        with tempfile.TemporaryDirectory() as tmp:
            mp, out = Path(tmp) / 'map.json', Path(tmp) / 'fit.json'
            record = {'page': 1, 'id': 5, 'name': 'sample', 'orig': original, 'new': target}
            if indices:
                record['paragraphs'] = indices
            mp.write_text(json.dumps([record]), encoding='utf-8')
            argv = ['ppt_fit.py', 'source.pptx', '--map', str(mp), '--out', str(out),
                    '--tol', '0', '--steps', steps]
            with patch.dict(sys.modules, {'win32com': win32, 'win32com.client': client}), \
                 patch.object(sys, 'argv', argv), patch.object(sys, 'path', [str(ROOT / 'scripts')] + sys.path):
                runpy.run_path(str(ROOT / 'scripts/ppt_fit.py'), run_name='__main__')
            return json.loads(out.read_text(encoding='utf-8'))[0]

    def test_every_paragraph_stays_at_or_above_original(self):
        result = self.run_fit([10, 20], [14, 22])
        self.assertTrue(all(v >= o for v, o in zip(result['applied'], [10, 20])))

    def test_blank_paragraphs_keep_physical_indices(self):
        result = self.run_fit([10, 20], [14, 22], indices=[1, 3])
        self.assertEqual(result['paragraphs'], [1, 3])
        self.assertTrue(all(v >= o for v, o in zip(result['applied'], [10, 20])))

    def test_floor_is_attempted_below_last_configured_step(self):
        result = self.run_fit([10], [20])
        self.assertEqual(result['applied'], [10])
        self.assertEqual(result['over'], 0)

    def test_measurement_failure_is_not_zero_overflow(self):
        with self.assertRaisesRegex(RuntimeError, '无法测量'):
            self.run_fit([10], [20], broken=True)


if __name__ == '__main__':
    unittest.main()
