"""Synthetic decks and fake COM prove collisions fail before any input is written."""
import json
import os
from pathlib import Path
import runpy
import subprocess
import sys
import tempfile
from types import ModuleType
import unittest
from unittest.mock import patch

from pptx import Presentation
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
SCRIPTS = ROOT/'scripts'


class OutputPathTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.src = self.root/'source.pptx'
        prs = Presentation()
        box = prs.slides.add_slide(prs.slide_layouts[6]).shapes.add_textbox(Inches(1), Inches(1), Inches(7), Inches(2))
        box.text = 'Synthetic source content'
        box.text_frame.paragraphs[0].runs[0].font.size = Pt(14)
        prs.save(self.src)
        self.original = self.src.read_bytes()
        self.fit = self.root/'fit.json'
        self.fit.write_text(json.dumps([{'page': 1, 'id': box.shape_id, 'paragraphs': [1], 'applied': [20]}]), encoding='utf-8')
        self.config = self.root/'config.json'
        self.config.write_text('{}', encoding='utf-8')

    def reject(self, script, *args):
        inputs = {p: p.read_bytes() for p in [self.src, self.fit, self.config]}
        proc = subprocess.run([sys.executable, '-X', 'utf8', str(SCRIPTS/script), *map(str, args)], capture_output=True)
        self.assertEqual(proc.returncode, 2, proc.stderr.decode('utf-8', 'replace'))
        for path, original in inputs.items():
            self.assertEqual(path.read_bytes(), original, str(path))

    def test_transforms_cannot_replace_source(self):
        for script, extra in [('ppt_enlarge.py', ['--map', self.root/'map.json']), ('ppt_white.py', []), ('ppt_deco.py', ['--config', self.config])]:
            with self.subTest(script=script):
                self.reject(script, self.src, self.src, *extra)

    def test_sidecar_cannot_replace_source_or_presentation(self):
        for output, mapping in [(self.root/'out.pptx', self.src), (self.root/'out.pptx', self.root/'out.pptx')]:
            with self.subTest(mapping=mapping):
                self.reject('ppt_enlarge.py', self.src, output, '--map', mapping)
        self.reject('ppt_white.py', self.src, self.root/'out.pptx', '--config', self.config, '--log', self.config)

    def test_apply_and_delivery_cannot_replace_inputs(self):
        self.reject('ppt_apply.py', self.src, '--fit', self.fit, '--out', self.src)
        self.reject('ppt_apply.py', self.src, '--fit', self.fit, '--out', self.root/'out.pptx', '--deliver', self.src)
        self.reject('ppt_apply.py', self.src, '--fit', self.fit, '--out', self.root/'out.pptx', '--deliver', self.fit)

    def test_measurement_rejects_collision_before_starting_office(self):
        win32, client = ModuleType('win32com'), ModuleType('win32com.client')
        def unexpected(*args):
            raise AssertionError('Office must not start for conflicting paths')
        client.DispatchEx = unexpected
        win32.client = client
        args = ['ppt_fit.py', str(self.src), '--map', str(self.fit), '--out', str(self.src)]
        with patch.dict(sys.modules, {'win32com': win32, 'win32com.client': client}), patch.object(sys, 'argv', args), patch.object(sys, 'path', [str(SCRIPTS)] + sys.path), self.assertRaises(SystemExit) as caught:
            runpy.run_path(str(SCRIPTS/'ppt_fit.py'), run_name='__main__')
        self.assertEqual(caught.exception.code, 2)
        self.assertEqual(self.src.read_bytes(), self.original)

    def test_report_and_scan_outputs_cannot_replace_source(self):
        for script in ('ppt_verify.py', 'ppt_audit.py'):
            with self.subTest(script=script):
                self.reject(script, self.src, '--json', self.src)
        for script in ('ppt_scan.py', 'ppt_chrome.py'):
            with self.subTest(script=script):
                self.reject(script, self.src, self.src)

    def test_existing_hardlink_alias_is_rejected(self):
        alias = self.root/'alias.pptx'
        try:
            os.link(self.src, alias)
        except OSError as error:
            self.skipTest(str(error))
        self.reject('ppt_white.py', self.src, alias)


if __name__ == '__main__':
    unittest.main()
