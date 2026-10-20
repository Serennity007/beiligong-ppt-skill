# -*- coding: utf-8 -*-
"""步骤 E-1：全篇字号分级放大 + 字色加深 + 收紧行距，并输出逐形状字号映射表

放大是"先大胆放大、再实测回退"的第一步。本脚本只负责放大，不做任何越框判断，
越框由 ppt_fit.py 实测后回退。两者必须配套使用。

用法：
    python ppt_enlarge.py 源文件.pptx 输出.pptx --map size_map.json [--config enlarge_config.json]

size_map.json 记录每个形状 [原字号, 目标字号]，是回退的地板依据（不允许缩到比原来还小）。
"""
import argparse
import json
import sys
from collections import Counter

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Pt

sys.path.insert(0, __file__.rsplit("\\", 1)[0].rsplit("/", 1)[0])
import pptkit as K

ap = argparse.ArgumentParser()
ap.add_argument("src")
ap.add_argument("out")
ap.add_argument("--map", required=True)
ap.add_argument("--config", default=None)
a = ap.parse_args()

DEF = {
    # [字号上界(不含), 倍率, 放大后下限]
    "curve": [[11, 1.40, 14.0], [14, 1.25, 16.0], [18, 1.20, 19.0],
              [22, 1.18, 23.0], [26, 1.13, 27.0], [999, 1.10, 34.0]],
    "default_pt": 18.0,
    "font_map": {
        "006633": "173404", "006C39": "173404", "044A2E": "122C03", "17865A": "173404",
        "595959": "2E2E2E", "6E6E6E": "3D3D3D", "222222": "141414", "404040": "262626",
        "4A4A4A": "2A2A2A", "9A9A9A": "555555", "BFBFBF": "666666", "C0C0C0": "666666",
        "CC3333": "B02020",
    },
    "fill_map": {"006633": "173404", "006C39": "173404"},
    "line_spacing": 1.0,
    "kill_space": True,
    "font": "微软雅黑",
    "skip_name": [],           # 形状名包含这些串的跳过（如页码）
    "skip_shape_pt": None,
}
cfg = K.load_config(a.config, DEF)
CURVE = cfg["curve"]
CMAP = {k.upper(): v.upper() for k, v in (cfg.get("font_map") or {}).items()}
FMAP = {k.upper(): v.upper() for k, v in (cfg.get("fill_map") or {}).items()}
SKIP = cfg["skip_name"] or []


def bump(pt):
    if pt is None:
        return float(cfg["default_pt"])
    for lim, fac, floor in CURVE:
        if pt < lim:
            return K.round_half(max(floor, pt * fac))
    return K.round_half(pt)


prs = Presentation(a.src)
mapping = []
stat = Counter()

for i, sl in enumerate(prs.slides, 1):
    for sh, _, _ in K.iter_shapes(sl.shapes, i):
        if any(s in (sh.name or "") for s in SKIP):
            continue
        g = K.geo(sh)
        try:
            f = sh.fill
            fill = K.fill_hex(sh)
            if fill in FMAP and g["w"] >= 0.4 and g["h"] >= 0.28:
                f.fore_color.rgb = RGBColor.from_string(FMAP[fill])
                stat["填充加深"] += 1
        except Exception:
            pass
        if not getattr(sh, "has_text_frame", False):
            continue
        tf = sh.text_frame
        pinfo = []
        for p in tf.paragraphs:
            rs = [r for r in p.runs if r.text.strip()]
            if not rs:
                continue
            cur = [(r.font.size.pt if r.font.size else None) for r in rs]
            base = max([c for c in cur if c] or [float(cfg["default_pt"])])
            pinfo.append({"runs": rs, "orig": base, "new": bump(base)})
        if not pinfo:
            continue
        stat["文字形状"] += 1
        for pi in pinfo:
            for r in pi["runs"]:
                r.font.size = Pt(pi["new"])
                stat["run"] += 1
                try:
                    c = str(r.font.color.rgb).upper()
                except Exception:
                    c = None
                if c in CMAP:
                    r.font.color.rgb = RGBColor.from_string(CMAP[c])
                    stat["字色加深"] += 1
                K.set_ea_font(r, cfg.get("font", K.FONT_DEFAULT))
        for p in tf.paragraphs:
            try:
                if p.line_spacing is not None and p.line_spacing > 1.05:
                    p.line_spacing = cfg["line_spacing"]
            except Exception:
                pass
            if cfg["kill_space"]:
                for attr in ("space_before", "space_after"):
                    try:
                        if getattr(p, attr):
                            setattr(p, attr, Pt(0))
                    except Exception:
                        pass
        mapping.append({"page": i, "id": sh.shape_id, "name": sh.name,
                        "orig": [pi["orig"] for pi in pinfo],
                        "new": [pi["new"] for pi in pinfo]})

prs.save(a.out)
json.dump(mapping, open(a.map, "w", encoding="utf-8"), ensure_ascii=False)
print("保存:", a.out)
print("统计:", dict(stat))
print("映射表:", a.map, len(mapping), "个形状")
