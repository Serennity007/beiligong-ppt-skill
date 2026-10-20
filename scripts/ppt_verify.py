# -*- coding: utf-8 -*-
"""步骤 F：交付前体检

查四件事：
  1. 还有没有"浅底白字"（投影上等于空白）；
  2. 还有没有该去掉的浅绿/浅灰底纹；
  3. 页数对不对、包有没有坏；
  4. 有没有越框（读 fit_result.json，可选）。

用法：
    python ppt_verify.py 成品.pptx [--defill E8F5E9,F1F8F1] [--fit fit_result.json] [--json 报告.json]
"""
import argparse
import json
import sys
from collections import Counter

from pptx import Presentation
from pptx.enum.dml import MSO_FILL_TYPE
from pptx.enum.shapes import MSO_SHAPE_TYPE

sys.path.insert(0, __file__.rsplit("\\", 1)[0].rsplit("/", 1)[0])
import pptkit as K

ap = argparse.ArgumentParser()
ap.add_argument("src")
ap.add_argument("--defill", default="E8F5E9,F1F8F1,EFF5EF,F5F5F5,FAFAFA")
ap.add_argument("--fit", default=None)
ap.add_argument("--tol", type=float, default=3.0)
ap.add_argument("--json", default=None)
a = ap.parse_args()

DEFILL = set(x.strip().upper() for x in a.defill.split(",") if x.strip())
LIGHT_TEXT = {"FFFFFF", "FFFEFE", "F5F5F5", "E8F5E9", "F2F2F2"}

prs = Presentation(a.src)
issues = []
stat = Counter()

for i, sl in enumerate(prs.slides, 1):
    try:
        f = sl.background.fill
        bg = str(f.fore_color.rgb).upper() if f.type == MSO_FILL_TYPE.SOLID else "FFFFFF"
    except Exception:
        bg = "FFFFFF"
    if bg != "FFFFFF" and K.luminance(bg) > 0.62:
        issues.append({"page": i, "type": "页面底色非白", "color": bg})
        stat["页面底色非白"] += 1
    # 先收集本页所有"实心深色块"和"图片"——它们是文字框真正压在上面的底
    dark_bands, covers = [], []
    for sh, _, _ in K.iter_shapes(sl.shapes, i):
        g0 = K.geo(sh)
        if sh.shape_type == MSO_SHAPE_TYPE.PICTURE:
            covers.append(g0)
            continue
        c0 = K.fill_hex(sh)
        if c0 and K.luminance(c0) <= 0.62:
            dark_bands.append((g0, c0))

    def behind(g):
        """文字框中心压在什么底色上；压在图片上返回 None（无法判定，跳过）"""
        cx, cy = g["l"] + g["w"] / 2, g["t"] + g["h"] / 2
        for cv in covers:
            if cv["l"] <= cx <= cv["l"] + cv["w"] and cv["t"] <= cy <= cv["t"] + cv["h"]:
                return None
        hit = None
        for gb, c0 in dark_bands:
            if gb["l"] <= cx <= gb["l"] + gb["w"] and gb["t"] <= cy <= gb["t"] + gb["h"]:
                if hit is None or gb["w"] * gb["h"] < hit[0]["w"] * hit[0]["h"]:
                    hit = (gb, c0)
        return hit[1] if hit else "FFFFFF"

    for sh, _, _ in K.iter_shapes(sl.shapes, i):
        g = K.geo(sh)
        fill = K.fill_hex(sh)
        has_light_text = False
        if getattr(sh, "has_text_frame", False):
            for p in sh.text_frame.paragraphs:
                for r in p.runs:
                    if not r.text.strip():
                        continue
                    try:
                        if str(r.font.color.rgb).upper() in LIGHT_TEXT:
                            has_light_text = True
                    except Exception:
                        pass
        if fill in DEFILL and g["w"] >= 0.4 and g["h"] >= 0.28:
            # 浅底本身不一定错：只要上面的字是深色，就是正常的浅色卡片
            t = "残留浅底(且有白字)" if has_light_text else "浅色卡片(字为深色,可保留)"
            if has_light_text:
                issues.append({"page": i, "type": t, "color": fill,
                               "name": sh.name, "text": K.text_of(sh)[:30]})
            stat[t] += 1
        if not getattr(sh, "has_text_frame", False):
            continue
        base = fill or behind(g)
        if base is None or K.luminance(base) <= 0.62:
            continue                       # 压在图片上不判；深底/彩底，白字没问题
        for p in sh.text_frame.paragraphs:
            for r in p.runs:
                if not r.text.strip():
                    continue
                try:
                    c = str(r.font.color.rgb).upper()
                except Exception:
                    continue
                if c in LIGHT_TEXT:
                    issues.append({"page": i, "type": "浅底白字", "name": sh.name,
                                   "text": r.text[:30]})
                    stat["浅底白字"] += 1

print("包内 sldId 数:", K.page_count_zip(a.src), " / python-pptx 页数:", len(prs.slides._sldIdLst))

if a.fit:
    fit = json.load(open(a.fit, encoding="utf-8"))
    over = [r for r in fit if r["over"] > a.tol]
    stat["越框"] = len(over)
    for r in over:
        issues.append({"page": r["page"], "type": "越框", "name": r["name"],
                       "over": r["over"], "applied": r["applied"]})
    print("越框框数:", len(over), "/", len(fit))

print("问题统计:", dict(stat))
for it in issues[:40]:
    print("  P%-4d %-8s %s" % (it["page"], it["type"], it.get("text") or it.get("color") or it.get("name", "")))
if len(issues) > 40:
    print("  ...还有", len(issues) - 40, "条")

if a.json:
    json.dump(issues, open(a.json, "w", encoding="utf-8"), ensure_ascii=False, indent=1)
    print("报告:", a.json)
