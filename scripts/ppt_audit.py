# -*- coding: utf-8 -*-
"""按《教师要求清单》逐条体检（skill 的验收环节）

能自动判的：白底、去底块、深字、院标完整、统摄标题连续、页内小标题、
右上角不放文字、底部横幅、编辑定位信息、案例三件套、字号下限、越框。
判不了的（总图画没画对、讲述顺序、口语化标题、连线压字）作为人工核对项列在报告末尾。

用法：
    python ppt_audit.py 成品.pptx [--json 报告.json] [--md 报告.md] [--fit fit_result.json]
          [--config audit_config.json]

退出码：0 = 无 FAIL（不代表人工验收完成）；1 = 有 FAIL
"""
import argparse
import json
import os
import re
import sys

from pptx import Presentation
from pptx.enum.dml import MSO_FILL_TYPE
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.oxml.ns import qn

sys.path.insert(0, __file__.rsplit("\\", 1)[0].rsplit("/", 1)[0])
import pptkit as K

ap = argparse.ArgumentParser()
ap.add_argument("src")
ap.add_argument("--json", default=None)
ap.add_argument("--md", default=None)
ap.add_argument("--fit", default=None, help="ppt_fit.py 的结果，用于查越框")
ap.add_argument("--tol", type=float, default=3.0)
ap.add_argument("--config", default=None)
a = ap.parse_args()
try:
    K.check_output_paths({'src': a.src, 'fit': a.fit, 'config': a.config}, {'json': a.json, 'md': a.md})
except (ValueError, OSError) as exc:
    ap.error(str(exc))

DEF = {
    "defill": ["E8F5E9", "F1F8F1", "EFF5EF", "F5F5F5", "FAFAFA"],
    "light_text": ["FFFFFF", "FFFEFE", "F5F5F5", "E8F5E9", "F2F2F2"],
    "lum_threshold": 0.62,
    "min_body_pt": 12.0,
    "logo": {"l_min": 10.5, "t_max": 1.0, "ratio_tol": 0.08, "skip_pages": [1]},
    "solo_title_ratio": 0.40,      # 只出现一页的统摄标题占比上限（总览/分隔页天然孤立）
    "bottom_ratio": 0.60,          # 有底部横幅的页占比上限
    "case_kw": ["案例", "事件", "复盘"],
    "case_problem_kw": ["问题"],
    "case_fix_kw": ["方案", "防控", "措施", "处置"],
    "locator_re": r"正文位置|本部分回答|本部分分几讲|本页讲什么|正文\s*P\s*[.．]?\s*\d|P\s*[.．]\s*\d+\s*[—\-~－]\s*\d+|页段|编辑说明|待补|TODO|E\d{2}编号",
}
cfg = K.load_config(a.config, DEF)
DEFILL = set(x.upper() for x in cfg["defill"])
LIGHT = set(x.upper() for x in cfg["light_text"])
TH = float(cfg["lum_threshold"])
LOC = re.compile(cfg["locator_re"])

prs = Presentation(a.src)
SW, SH = prs.slide_width / K.IN, prs.slide_height / K.IN
pages = []

for i, sl in enumerate(prs.slides, 1):
    try:
        f = sl.background.fill
        bg = str(f.fore_color.rgb).upper() if f.type == MSO_FILL_TYPE.SOLID else "FFFFFF"
    except Exception:
        bg = "FFFFFF"
    rec = {"page": i, "bg": bg, "title": "", "sub": "", "bottom": False,
           "pics": [], "texts": [], "min_pt": None, "tr_texts": [], "bands": []}
    for sh, _, _ in K.iter_shapes(sl.shapes, i):
        g = K.geo(sh)
        # 纯色块（含无文字的色带）也要记，否则会把"绿带上的白字"误判成浅底白字
        _f = K.fill_hex(sh)
        if _f and K.luminance(_f) <= TH:
            rec["bands"].append({"geo": g, "fill": _f})
        if sh.shape_type == MSO_SHAPE_TYPE.PICTURE:
            nat = None
            try:
                w0, h0 = sh.image.size
                nat = (w0 / h0) if h0 else None
            except Exception:
                pass
            rect = sh._element.find(qn("p:blipFill"))
            crop = rect.find(qn("a:srcRect")) if rect is not None else None
            has_src = crop is not None and any(int(crop.get(edge, "0")) != 0
                                               for edge in ("l", "t", "r", "b"))
            rec["pics"].append({"l": g["l"], "t": g["t"], "w": g["w"], "h": g["h"],
                                "name": sh.name, "srcRect": has_src,
                                "ratio": (g["w"] / g["h"]) if g["h"] else None,
                                "nat": nat})
        txt = K.text_of(sh)
        if not txt.strip():
            continue
        fill = K.fill_hex(sh)
        runs = []
        for p in sh.text_frame.paragraphs:
            for r in p.runs:
                if not r.text.strip():
                    continue
                try:
                    c = str(r.font.color.rgb).upper()
                except Exception:
                    c = None
                pt = r.font.size.pt if r.font.size else None
                runs.append({"text": r.text, "color": c, "pt": pt})
        rec["texts"].append({"geo": g, "fill": fill, "text": txt, "runs": runs, "name": sh.name})
        if g["t"] < 0.75 and g["l"] < 5.0 and g["w"] > 3.0 and not rec["title"]:
            rec["title"] = txt.strip()
        if 0.88 <= g["t"] <= 1.25 and g["l"] < 1.6 and g["w"] > 8.0 and not rec["sub"]:
            rec["sub"] = txt.strip()
        if g["t"] > SH - 1.9 and g["w"] > 8.0:
            rec["bottom"] = True
        if g["l"] > 9.5 and g["t"] < 0.9:
            rec["tr_texts"].append(txt.strip()[:40])
    for t in rec["texts"]:
        for r in t["runs"]:
            if r["pt"]:
                rec["min_pt"] = r["pt"] if rec["min_pt"] is None else min(rec["min_pt"], r["pt"])
    pages.append(rec)

n = len(pages)
results = []


def add(cid, name, level, detail, items=None):
    results.append({"id": cid, "name": name, "level": level,
                    "detail": detail, "items": (items or [])[:30]})


# ---------- E1 白底 ----------
bad = [p["page"] for p in pages if p["bg"] != "FFFFFF"]
add("E1", "页面白底", "FAIL" if bad else "PASS",
    "非白底 %d 页" % len(bad), [{"page": x, "bg": pages[x - 1]["bg"]} for x in bad])

# ---------- E1b 去底块 ----------
bad, cards = [], 0
for p in pages:
    for t in p["texts"]:
        if t["fill"] in DEFILL and t["geo"]["w"] >= 0.4 and t["geo"]["h"] >= 0.28:
            bad.append({"page": p["page"], "text": t["text"][:30], "fill": t["fill"]})
add("E1b", "无浅色铺底", "FAIL" if bad else "PASS",
    "带文字的浅色铺底 %d 处；无文字底块、组合和位图另行目检" % len(bad), bad)

# ---------- E2 深字（看背后压的是什么颜色） ----------
bad = []
for p in pages:
    dark = [(b["geo"], b["fill"]) for b in p["bands"]]
    covers = [{"l": q["l"], "t": q["t"], "w": q["w"], "h": q["h"]} for q in p["pics"]]
    for t in p["texts"]:
        base = t["fill"]
        if not base:
            cx = t["geo"]["l"] + t["geo"]["w"] / 2
            cy = t["geo"]["t"] + t["geo"]["h"] / 2
            if any(cv["l"] <= cx <= cv["l"] + cv["w"] and cv["t"] <= cy <= cv["t"] + cv["h"] for cv in covers):
                continue                      # 压在图片上，不判
            hit = None
            for gb, c0 in dark:
                if gb["l"] <= cx <= gb["l"] + gb["w"] and gb["t"] <= cy <= gb["t"] + gb["h"]:
                    if hit is None or gb["w"] * gb["h"] < hit[0]["w"] * hit[0]["h"]:
                        hit = (gb, c0)
            base = hit[1] if hit else "FFFFFF"
        if K.luminance(base) <= TH:
            continue
        for r in t["runs"]:
            if r["color"] in LIGHT:
                bad.append({"page": p["page"], "text": r["text"][:30]})
                break
add("E2", "浅底无白字", "FAIL" if bad else "PASS", "浅底白字 %d 处" % len(bad), bad)

# ---------- E6 院标完整 ----------
lg = cfg["logo"]
skip = set(lg.get("skip_pages", [1]))
miss, cropped, distort = [], [], []
for p in pages:
    # 跳过封面/目录/致谢等无统摄标题的页（用 skip_pages 显式指定，或该页根本没有标题）
    if p["page"] in skip or not p["title"]:
        continue
    cand = [q for q in p["pics"] if q["l"] > lg["l_min"] and q["t"] < lg["t_max"]]
    if not cand:
        miss.append({"page": p["page"]})
        continue
    for q in cand:
        if q["srcRect"]:
            cropped.append({"page": p["page"], "name": q["name"]})
        if q["ratio"] and q["nat"]:
            if abs(q["ratio"] - q["nat"]) / q["nat"] > lg["ratio_tol"]:
                distort.append({"page": p["page"], "ratio": round(q["ratio"], 3),
                                "nat": round(q["nat"], 3), "name": q["name"]})
lv = "FAIL" if (miss or cropped) else ("WARN" if distort else "PASS")
add("E6", "院标完整未裁切", lv,
    "缺 %d 页 / 被裁 %d / 比例失真 %d" % (len(miss), len(cropped), len(distort)),
    miss + cropped + distort)

# ---------- C1/C2 统摄标题连续 ----------
seq = [(p["page"], p["title"]) for p in pages if p["title"]]
groups, cur = [], None
for pg, t in seq:
    if cur and cur["t"] == t:
        cur["pages"].append(pg)
    else:
        cur = {"t": t, "pages": [pg]}
        groups.append(cur)
solo = [g["pages"][0] for g in groups if len(g["pages"]) == 1]
ratio_solo = len(solo) / max(1, len(groups))
lv = "PASS" if ratio_solo <= cfg["solo_title_ratio"] else "WARN"
add("C1", "统摄标题成组连续", lv,
    "%d 组 / %d 页；孤立页 %d（%.0f%%）" % (len(groups), len(seq), len(solo), ratio_solo * 100),
    [{"page": x, "title": pages[x - 1]["title"][:40]} for x in solo])

# ---------- C3 页内小标题 ----------
# 孤立页（总览 / 示例 / 部分收束）本就不设小标题，跳过
miss = [p["page"] for p in pages
        if p["title"] and not p["sub"] and p["page"] not in set(solo)]
lv = "PASS" if not miss else ("WARN" if len(miss) < n * 0.2 else "FAIL")
add("C3", "页内小标题齐备", lv, "缺小标题 %d 页" % len(miss), [{"page": x} for x in miss])

# ---------- A4 右上角不放文字 ----------
bad = []
for p in pages:
    for t in p["tr_texts"]:
        if re.fullmatch(r"\d{1,3}", t or ""):
            continue                       # 页码可以留
        bad.append({"page": p["page"], "text": t})
add("A4", "右上角不放文字标题", "WARN" if bad else "PASS", "%d 处" % len(bad), bad)

# ---------- E4 底部横幅 ----------
cnt = sum(1 for p in pages if p["bottom"])
r = cnt / max(1, n)
lv = "PASS" if r <= cfg["bottom_ratio"] else "WARN"
add("E4", "底部横幅不逐页补", lv, "%d/%d 页有底部条（%.0f%%）" % (cnt, n, r * 100), [])

# ---------- B4/B5/F1 编辑定位信息 ----------
bad = []
for p in pages:
    for t in p["texts"]:
        m = LOC.search(t["text"])
        if m:
            bad.append({"page": p["page"], "text": t["text"][:40], "hit": m.group(0)})
add("F1", "无编辑定位信息", "FAIL" if bad else "PASS", "%d 处" % len(bad), bad)

# ---------- D1/D2 案例三件套 ----------
bad = []
for p in pages:
    ti = p["title"]
    if not any(k in ti for k in cfg["case_kw"]):
        continue
    alltxt = ti + " " + " ".join(t["text"] for t in p["texts"])
    if len(alltxt) < 40:
        continue                       # 字太少，是总览/分隔页，不是案例展开页
    has_p = any(k in alltxt for k in cfg["case_problem_kw"])
    has_f = any(k in alltxt for k in cfg["case_fix_kw"])
    if not (has_p and has_f):
        bad.append({"page": p["page"], "title": ti[:40], "问题": has_p, "方案": has_f})
add("D1", "案例讲清问题与方案", "WARN" if bad else "PASS", "缺件 %d 页" % len(bad), bad)

# ---------- F2 字号下限 ----------
bad = []
for p in pages:
    for t in p["texts"]:
        for r in t["runs"]:
            if r["pt"] and r["pt"] < cfg["min_body_pt"] and len(t["text"]) > 12:
                bad.append({"page": p["page"], "pt": r["pt"], "text": r["text"][:24]})
lv = "FAIL" if bad else "PASS"
add("F2", "正文字号不低于 %.0fpt" % cfg["min_body_pt"], lv, "小字 %d 处" % len(bad), bad)

# ---------- F3 越框（可选） ----------
if a.fit:
    fit = json.load(open(a.fit, encoding="utf-8"))
    over = [r for r in fit if r["over"] > a.tol]
    add("F3", "无文字越框", "PASS" if not over else "FAIL", "%d 框越框" % len(over),
        [{"page": r["page"], "name": r["name"], "over": r["over"]} for r in over])
    if not fit:
        add("F3-data", "实测数据有效", "FAIL", "实测结果为空，不能作为越框通过的证据")
else:
    add("F3", "无文字越框", "WARN", "未提供 --fit，未执行实测检查；不能宣称零越框")

# ---------- 输出 ----------
fails = [r for r in results if r["level"] == "FAIL"]
warns = [r for r in results if r["level"] == "WARN"]
print("体检：%s（%d 页）" % (a.src, n))
for r in results:
    print("  [%-4s] %-22s %s" % (r["level"], r["name"], r["detail"]))
print("合计：FAIL %d / WARN %d / PASS %d" % (len(fails), len(warns),
                                             sum(1 for r in results if r["level"] == "PASS")))

MANUAL = [
    "B2/B3 总图画的是组成与关系，不是把标题列表塞进方框（箭头须有明确含义）",
    "B1 目录只呈现主题与顺序，无解释性导读",
    "C4 统摄标题具体（对象+专业概念），不是同一个空泛标题写全篇",
    "C5 标题无口语、制作说明、'几个方面''一些思考'、同义反复",
    "D3 案例结论前置，不是先堆时间线/来源/评分",
    "D4 实际处置与防控建议分开写，不把建议说成对方已实施",
    "E3 顶部横幅白底+深色标题+细线，无大面积铺色",
    "E5 图表内白字深底已同步处理；位图固化的底块已重绘",
    "E7 实景照片与院标保留原色",
    "F3 逐页目检：贴边、连线压字、图注遮挡、过密小字",
    "G  交付八问（结构/连贯性/专业表述/投影可读性）",
]

if a.json:
    json.dump({"src": a.src, "n": n, "checks": results, "manual": MANUAL},
              open(a.json, "w", encoding="utf-8"), ensure_ascii=False, indent=1)
    print("JSON:", a.json)

if a.md:
    L = ["# 教师要求体检报告", "", "文件：`%s`，共 %d 页" % (os.path.basename(a.src), n), "",
         "| 结果 | 编号 | 项目 | 结论 |", "|---|---|---|---|"]
    for r in results:
        L.append("| %s | %s | %s | %s |" % (r["level"], r["id"], r["name"], r["detail"]))
    L += ["", "## 明细"]
    for r in results:
        if not r["items"]:
            continue
        L += ["", "### %s %s（%d 条）" % (r["id"], r["name"], len(r["items"]))]
        for it in r["items"]:
            L.append("- P%s %s" % (it.get("page"), " ".join(
                "%s=%s" % (k, v) for k, v in it.items() if k != "page")))
    L += ["", "## 人工核对项（脚本判不了，逐条勾）", ""]
    L += ["- [ ] " + m for m in MANUAL]
    open(a.md, "w", encoding="utf-8").write("\n".join(L))
    print("Markdown:", a.md)

sys.exit(1 if fails else 0)
