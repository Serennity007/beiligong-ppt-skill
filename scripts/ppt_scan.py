# -*- coding: utf-8 -*-
"""扫描 pptx 版式骨架：形状、位置、填充色、文字颜色 → JSON

用法:
    python ppt_scan.py 输入.pptx 输出.json
"""
import argparse, json
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


def inches(v):
    return None if v is None else round(v / 914400.0, 3)


def fill_info(sh):
    try:
        f = sh.fill
        t = str(f.type)
        info = {"type": t}
        if t.startswith("SOLID"):
            try:
                info["fg"] = str(f.fore_color.rgb)
            except Exception:
                pass
        return info
    except Exception:
        return None


def run_colors(tf):
    out = []
    try:
        for p in tf.paragraphs:
            for r in p.runs:
                c = None
                try:
                    c = str(r.font.color.rgb)
                except Exception:
                    pass
                out.append({"t": r.text[:40], "color": c,
                            "sz": (r.font.size.pt if r.font.size else None),
                            "b": r.font.bold})
    except Exception:
        pass
    return out


def walk(shapes):
    res = []
    for sh in shapes:
        d = {"name": str(sh.name), "type": str(sh.shape_type), "id": sh.shape_id,
             "l": inches(sh.left), "t": inches(sh.top),
             "w": inches(sh.width), "h": inches(sh.height),
             "fill": fill_info(sh)}
        try:
            d["line"] = (str(sh.line.color.rgb)
                         if (sh.line.fill.type is not None
                             and str(sh.line.fill.type).startswith("SOLID")) else None)
        except Exception:
            d["line"] = None
        if sh.has_text_frame:
            d["text"] = sh.text_frame.text[:200]
            d["runs"] = run_colors(sh.text_frame)
        if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
            d["children"] = walk(sh.shapes)
        res.append(d)
    return res


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("src")
    ap.add_argument("out")
    a = ap.parse_args()
    prs = Presentation(a.src)
    data = {"src": a.src,
            "slide_w": inches(prs.slide_width), "slide_h": inches(prs.slide_height),
            "n": len(prs.slides._sldIdLst), "slides": []}
    for i, s in enumerate(prs.slides, 1):
        data["slides"].append({"idx": i,
                               "layout": s.slide_layout.name if s.slide_layout else None,
                               "shapes": walk(s.shapes)})
    json.dump(data, open(a.out, "w", encoding="utf-8"), ensure_ascii=False, indent=1)
    print("slides:", data["n"], "size:", data["slide_w"], "x", data["slide_h"], "->", a.out)


if __name__ == "__main__":
    main()
