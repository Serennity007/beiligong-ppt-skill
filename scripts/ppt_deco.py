# -*- coding: utf-8 -*-
"""步骤 D：版式装饰 —— 页内小标题绿色横带、右上角院标、目录页重排

三个模块都靠 config 开关，互不影响，可单独跑。

用法：
    python ppt_deco.py 源文件.pptx 输出.pptx --config deco_config.json

deco_config.json 关键字段（详见 ../case/deco_config.json）：
  band:   {"on":true,"l":0.40,"t":0.96,"w":12.53,"h":0.44,"fill":"006633",
           "text_fill":"FFFFFF","size":15,"skip":[1,2],
           "match":{"t0":0.88,"t1":1.18,"l_max":1.6,"w_min":9}}
  logo:   {"on":true,"source_slide":1,"h":0.56,"right":0.38,"top":0.14,"band_h":0.73,
           "name":"v59b:logo","skip":[1],"match":{"t_max":1.0,"l_min":10.5,"w_max":2.5,"h_max":1.0}}
  toc:    {"on":true,"page":2,"rows":[["01","标题","副标"], ...]}
"""
import argparse
import copy
import json
import sys
from collections import Counter

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE_TYPE, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

sys.path.insert(0, __file__.rsplit("\\", 1)[0].rsplit("/", 1)[0])
import pptkit as K

ap = argparse.ArgumentParser()
ap.add_argument("src")
ap.add_argument("out")
ap.add_argument("--config", required=True)
a = ap.parse_args()

cfg = K.load_config(a.config)
FONT = cfg.get("font", K.FONT_DEFAULT)
prs = Presentation(a.src)
slides = list(prs.slides)
n = len(slides)
log = []


def px(v):
    return Emu(int(v * K.IN))


def add_text(slide, l, t, w, h, lines, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(px(l), px(t), px(w), px(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, (txt, size, bold, color) in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        r = p.add_run()
        r.text = txt
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.color.rgb = RGBColor.from_string(color.upper())
        K.set_ea_font(r, FONT)
    return box


# ============ 模块 1：页内小标题绿色横带 ============
band = cfg.get("band") or {}
if band.get("on"):
    B = band
    m = B.get("match", {})
    skip = set(B.get("skip", []))
    cnt = 0
    for idx, sl in enumerate(slides, 1):
        if idx in skip:
            continue
        target = None
        for sh in sl.shapes:
            if sh.shape_type != MSO_SHAPE_TYPE.TEXT_BOX or not sh.has_text_frame:
                continue
            g = K.geo(sh)
            txt = K.text_of(sh).strip()
            if not txt:
                continue
            if (m.get("t0", 0.88) <= g["t"] <= m.get("t1", 1.18)
                    and g["l"] < m.get("l_max", 1.6) and g["w"] > m.get("w_min", 9)):
                target = sh
                break
        if target is None:
            continue
        rect = sl.shapes.add_shape(MSO_SHAPE.RECTANGLE,
                                   px(B["l"]), px(B["t"]), px(B["w"]), px(B["h"]))
        rect.fill.solid()
        rect.fill.fore_color.rgb = RGBColor.from_string(B.get("fill", "006633").upper())
        rect.line.fill.background()
        rect.shadow.inherit = False
        target._element.addprevious(rect._element)   # 带子在文字下面
        target.left, target.top = px(B["l"] + 0.28), px(B["t"])
        target.width, target.height = px(B["w"] - 0.53), px(B["h"])
        tf = target.text_frame
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        for p in tf.paragraphs:
            for r in p.runs:
                r.font.size = Pt(B.get("size", 15))
                r.font.bold = True
                r.font.color.rgb = RGBColor.from_string(B.get("text_fill", "FFFFFF").upper())
                K.set_ea_font(r, FONT)
        cnt += 1
        log.append({"page": idx, "what": "小标题横带", "text": K.text_of(target)[:24]})
    print("小标题横带:", cnt)

# ============ 模块 2：右上角院标 ============
lg = cfg.get("logo") or {}
if lg.get("on"):
    src_idx = int(lg.get("source_slide", 1))
    src_slide = slides[src_idx - 1]
    tpl, img_part = None, None
    want = lg.get("source_name")
    for sh in src_slide.shapes:
        if sh.shape_type == MSO_SHAPE_TYPE.PICTURE:
            if want and sh.name != want:
                continue
            tpl = copy.deepcopy(sh._element)
            blip = sh._element.blipFill.find(qn("a:blip"))
            img_part = src_slide.part.related_part(blip.get(qn("r:embed")))
            break
    if tpl is None:
        raise SystemExit("源页未找到院标图片")
    if tpl.find(qn("p:spPr")).find(qn("a:srcRect")) is not None:
        print("警告：源图带 srcRect（可能被裁过），投影仍是缺角，请换一张未裁剪的源图")

    H = float(lg.get("h", 0.56))
    ratio = lg.get("ratio")  # 宽/高；不给就从模板读
    if not ratio:
        ext = tpl.find(qn("p:spPr")).find(qn("a:xfrm")).find(qn("a:ext"))
        ratio = int(ext.get("cx")) / int(ext.get("cy"))
    W = H * float(ratio)
    L = 13.333 - float(lg.get("right", 0.38)) - W
    bh = float(lg.get("band_h", 0.73))
    T = float(lg.get("top", 0.14)) + (bh - H) / 2
    m = lg.get("match", {})
    skip = set(lg.get("skip", []))
    nid = int(lg.get("id_base", 99000))
    cnt = 0

    def place(sl, idx):
        global nid
        for sh in list(sl.shapes):
            if sh.shape_type == MSO_SHAPE_TYPE.PICTURE:
                g = K.geo(sh)
                if (g["t"] < m.get("t_max", 1.0) and g["l"] > m.get("l_min", 10.5)
                        and g["w"] < m.get("w_max", 2.5) and g["h"] < m.get("h_max", 1.0)):
                    K.remove_element(sh)
                    log.append({"page": idx, "what": "删旧院标", "name": sh.name})
        el = copy.deepcopy(tpl)
        blip = el.blipFill.find(qn("a:blip"))
        rid = sl.part.relate_to(img_part, RT.IMAGE)
        blip.set(qn("r:embed"), rid)
        cnv = el.find(qn("p:nvPicPr")).find(qn("p:cNvPr"))
        nid += 1
        cnv.set("id", str(nid))
        cnv.set("name", lg.get("name", "v59b:logo"))
        spPr = el.find(qn("p:spPr"))
        xfrm = spPr.find(qn("a:xfrm"))
        if xfrm is None:
            xfrm = spPr.makeelement(qn("a:xfrm"), {})
            spPr.append(xfrm)
        for tag in ("a:off", "a:ext"):
            if xfrm.find(qn(tag)) is None:
                xfrm.append(xfrm.makeelement(qn(tag), {}))
        off, ext = xfrm.find(qn("a:off")), xfrm.find(qn("a:ext"))
        off.set("x", str(int(L * K.IN)))
        off.set("y", str(int(T * K.IN)))
        ext.set("cx", str(int(W * K.IN)))
        ext.set("cy", str(int(H * K.IN)))
        sl.shapes._spTree.append(el)
        log.append({"page": idx, "what": "加院标", "size": "%.2fx%.2f" % (W, H)})

    for idx, sl in enumerate(slides, 1):
        if idx in skip:
            continue
        place(sl, idx)
        cnt += 1
    print("院标放置:", cnt, "尺寸 %.2f x %.2f 英寸" % (W, H))

# ============ 模块 3：目录页重排 ============
toc = cfg.get("toc") or {}
if toc.get("on"):
    p = int(toc.get("page", 2))
    sl = slides[p - 1]
    drop_kw = toc.get("drop_text") or []
    for sh in list(sl.shapes):
        g = K.geo(sh)
        txt = K.text_of(sh).strip()
        drop = False
        if sh.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE and g["w"] > 12 and g["t"] < 0.5:
            drop = True
        elif sh.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE and 0.15 < g["w"] < 0.4 and g["t"] < 0.5:
            drop = True
        elif txt and (txt in drop_kw or any(k in txt for k in drop_kw)):
            drop = True
        if drop:
            K.remove_element(sh)
            log.append({"page": p, "what": "目录重排-删", "text": txt[:20]})
    L = toc.get("layout", {})
    add_text(sl, L.get("label_l", 1.15), L.get("label_t", 2.62), 3.0, 1.0,
             [(toc.get("label", "目  录"), L.get("label_size", 44), True, toc.get("label_color", "173404"))])
    add_text(sl, L.get("label_l", 1.15) + 0.03, L.get("label_t", 2.62) + 0.96, 3.0, 0.45,
             [(toc.get("label_en", "CONTENT"), L.get("en_size", 15), True, "6E6E6E")], anchor=MSO_ANCHOR.TOP)
    bar = sl.shapes.add_shape(MSO_SHAPE.RECTANGLE,
                              px(L.get("bar_l", 0.92)), px(L.get("bar_t", 2.70)),
                              px(L.get("bar_w", 0.11)), px(L.get("bar_h", 1.22)))
    bar.fill.solid()
    bar.fill.fore_color.rgb = RGBColor.from_string(toc.get("bar_color", "006633").upper())
    bar.line.fill.background()
    bar.shadow.inherit = False
    y = L.get("row_t", 1.62)
    for num, title, subs in toc.get("rows", []):
        chip = sl.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                   px(L.get("num_l", 4.45)), px(y + 0.04),
                                   px(0.62), px(L.get("chip_h", 0.40)))
        chip.adjustments[0] = 0.25
        chip.fill.solid()
        chip.fill.fore_color.rgb = RGBColor.from_string(toc.get("bar_color", "006633").upper())
        chip.line.fill.background()
        chip.shadow.inherit = False
        tf = chip.text_frame
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        pr = tf.paragraphs[0]
        pr.alignment = PP_ALIGN.CENTER
        r = pr.add_run()
        r.text = num
        r.font.size = Pt(L.get("num_size", 15))
        r.font.bold = True
        r.font.color.rgb = RGBColor.from_string("FFFFFF")
        K.set_ea_font(r, FONT)
        add_text(sl, L.get("title_l", 5.32), y, 7.5, 0.42,
                 [(title, L.get("title_size", 20), True, toc.get("label_color", "173404"))])
        add_text(sl, L.get("title_l", 5.32), y + 0.44, 7.6, 0.34,
                 [(subs, L.get("sub_size", 12), False, "4A4A4A")])
        y += L.get("row_gap", 0.86)
    log.append({"page": p, "what": "目录重排-建", "rows": len(toc.get("rows", []))})
    print("目录页重排完成，条目:", len(toc.get("rows", [])))

prs.save(a.out)
print("保存:", a.out)
print("改动统计:", dict(Counter(x["what"] for x in log)))
