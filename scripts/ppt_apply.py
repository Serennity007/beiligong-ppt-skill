# -*- coding: utf-8 -*-
"""步骤 E-3：把实测得到的最终字号写回 PPTX（用 python-pptx 保存，不让 COM 改写文件）

COM 只用来"量"，写回一律走 python-pptx —— 避免 PowerPoint 顺手把稿子另存一遍
（那会丢掉我们用 XML 直接塞进去的院标等元素）。

用法：
    python ppt_apply.py 上一步.pptx --fit fit_result.json --out 成品.pptx [--deliver 交付文件.pptx]
"""
import argparse
import json
import os
import shutil
import sys
from collections import Counter

from pptx import Presentation
from pptx.util import Pt

sys.path.insert(0, __file__.rsplit("\\", 1)[0].rsplit("/", 1)[0])
import pptkit as K

ap = argparse.ArgumentParser()
ap.add_argument("src")
ap.add_argument("--fit", required=True)
ap.add_argument("--out", required=True)
ap.add_argument("--deliver", default=None, help="写回后再覆盖到这个交付文件")
a = ap.parse_args()

fit = json.load(open(a.fit, encoding="utf-8"))
mp = {(r["page"], r["id"]): r for r in fit}
if not fit or len(mp) != len(fit):
    raise SystemExit("实测结果为空或包含重复形状，拒绝写回")

prs = Presentation(a.src)
hit = 0
sizes = Counter()

for i, sl in enumerate(prs.slides, 1):
    for sh, _, _ in K.iter_shapes(sl.shapes, i):
        if not getattr(sh, "has_text_frame", False):
            continue
        record = mp.get((i, sh.shape_id))
        if record is None:
            continue
        vals = record["applied"]
        paras = sh.text_frame.paragraphs
        indices = record.get("paragraphs", [idx for idx, p in enumerate(paras, 1)
                                             if any(r.text.strip() for r in p.runs)])
        if (len(indices) != len(vals) or len(set(indices)) != len(indices)
                or any(not isinstance(idx, int) or idx < 1 or idx > len(paras) for idx in indices)
                or any(v <= 0 for v in vals)):
            raise SystemExit(f"P{i} {sh.name}: 段落与实测结果不匹配，拒绝写回")
        for idx, v in zip(indices, vals):
            p = paras[idx - 1]
            rs = [r for r in p.runs if r.text.strip()]
            if not rs:
                raise SystemExit(f"P{i} {sh.name}: 实测段落已无文字，拒绝写回")
            for r in rs:
                r.font.size = Pt(v)
                sizes[v] += 1
        hit += 1

if hit != len(fit):
    raise SystemExit(f"仅匹配 {hit}/{len(fit)} 个形状，拒绝写回不完整结果")
prs.save(a.out)
print("写回形状:", hit, "/", len(fit))
print("包内 sldId 数:", K.page_count_zip(a.out), "（与源一致才安全）")
for k, v in sizes.most_common(14):
    print("   %-5s : %d" % (k, v))

if a.deliver:
    if not os.path.exists(a.deliver):
        raise SystemExit("交付文件不存在，拒绝自动创建：" + a.deliver)
    shutil.copyfile(a.out, a.deliver)
    print("交付文件已更新:", a.deliver)
